"""Download new course files: summarize them for the digest and optionally keep a local copy."""

from __future__ import annotations

import hashlib
import io
import json
import logging
import re
from pathlib import Path

from .collect import Course, Digest, FileItem, Module, html_to_text
from .config import Config
from .moodle import Moodle

log = logging.getLogger(__name__)

SUMMARIZABLE = {".pdf", ".pptx", ".docx", ".txt", ".md", ".html"}
MAX_SUMMARY_FILES = 8  # per run, to stay well inside the AI free tier's limits
EXCERPT_CHARS = 20_000  # per file, for the text fallback; the prompt says these are excerpts
MAX_DOWNLOAD_BYTES = 100 * 1024 * 1024
MAX_AI_PDF_BYTES = 30 * 1024 * 1024


def extract_text(data: bytes, filename: str) -> str:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    if ext == ".pptx":
        from pptx import Presentation

        slides = []
        for i, slide in enumerate(Presentation(io.BytesIO(data)).slides, 1):
            parts = [s.text_frame.text for s in slide.shapes if s.has_text_frame and s.text_frame.text]
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text:
                parts.append("Notes: " + slide.notes_slide.notes_text_frame.text)
            slides.append(f"[Slide {i}] " + "\n".join(parts))
        return "\n\n".join(slides)
    if ext == ".docx":
        import docx

        return "\n".join(p.text for p in docx.Document(io.BytesIO(data)).paragraphs if p.text)
    if ext in (".txt", ".md"):
        return data.decode("utf-8", errors="replace")
    if ext == ".html":
        return html_to_text(data.decode("utf-8", errors="replace"))
    return ""


RECORDING_TYPES = {".mp4", ".mov", ".m4v", ".avi", ".mkv", ".webm", ".wmv", ".flv", ".mpg", ".mpeg",
                   ".mp3", ".m4a", ".wav", ".aac", ".ogg", ".wma", ".flac"}


def _safe(name: str) -> str:
    return re.sub(r'[<>:"/\\|?*\x00-\x1f]+', "_", name).strip(" .") or "_"


def is_recording(f: FileItem) -> bool:
    return (Path(f.name).suffix.lower() in RECORDING_TYPES
            or f.mimetype.startswith(("video/", "audio/")))


def download_path(cfg: Config, course: Course, module: Module, f: FileItem) -> Path:
    """<DOWNLOAD_DIR>/<CSSD 3131 - Machine Learning>/<unit>/<week>/[<folder>/]<file>"""
    parts = [f"{course.shortname} - {course.title}"]
    parts += [p.strip() for p in module.section.split("›") if p.strip()]  # nested sections
    if module.modname == "folder":
        parts.append(module.name)
    parts += [p for p in f.name.split("/") if p]
    return cfg.download_dir.joinpath(*[_safe(p) for p in parts])


class SummaryCache:
    """File summaries already made, keyed by file and version, so a re-run never redoes them."""

    def __init__(self, path: Path):
        self.path = path
        try:
            self.data = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            self.data = {}

    @staticmethod
    def key(f: FileItem) -> str:
        return f"{f.url}@{f.timemodified}"

    def get(self, f: FileItem) -> str:
        return self.data.get(self.key(f), "")

    def put(self, f: FileItem, summary: str) -> None:
        self.data[self.key(f)] = summary
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self.path.write_text(json.dumps(self.data), encoding="utf-8")


class DownloadManifest:
    """What's already been downloaded, kept in <DOWNLOAD_DIR>/.eclass-downloads.json.

    Each file is keyed by its eClass item and name (stable across re-uploads) and remembers the
    version it was downloaded at, so the 2-hourly checks and the evening digest never fetch the
    same thing twice, and a copy posted in two places is only stored once."""

    def __init__(self, cfg: Config):
        self.root = cfg.download_dir
        self.path = self.root / ".eclass-downloads.json"
        try:
            self.data: dict = json.loads(self.path.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            self.data = {}

    @staticmethod
    def key(module: Module, f: FileItem) -> str:
        return f"{module.cmid}/{f.name}"

    def local_path(self, module: Module, f: FileItem) -> Path | None:
        entry = self.data.get(self.key(module, f))
        path = self.root / entry["path"] if entry else None
        return path if path and path.exists() else None

    def is_current(self, module: Module, f: FileItem) -> bool:
        entry = self.data.get(self.key(module, f))
        return bool(entry) and entry["version"] == f.timemodified and self.local_path(module, f) is not None

    def record(self, course: Course, module: Module, f: FileItem, path: Path, sha1: str = "",
               duplicate_of: str = "") -> None:
        self.data[self.key(module, f)] = {
            "course": course.id, "name": Path(f.name).name, "size": f.size, "version": f.timemodified,
            "path": path.relative_to(self.root).as_posix(), "sha1": sha1,
            **({"duplicate_of": duplicate_of} if duplicate_of else {})}

    def same_file_elsewhere(self, course: Course, module: Module, f: FileItem) -> str | None:
        """Another item in this course with the same file name and size (a re-post)."""
        me = self.key(module, f)
        for key, e in self.data.items():
            if (key != me and e["course"] == course.id and e["name"] == Path(f.name).name
                    and e["size"] == f.size and (self.root / e["path"]).exists()):
                return key
        return None

    def same_content(self, sha1: str, exclude: str) -> str | None:
        for key, e in self.data.items():
            if key != exclude and e.get("sha1") == sha1 and (self.root / e["path"]).exists():
                return key
        return None

    def save(self) -> None:
        self.root.mkdir(parents=True, exist_ok=True)
        self.path.write_text(json.dumps(self.data, indent=1), encoding="utf-8")


def _plan_download(manifest: DownloadManifest, cfg: Config, course: Course, module: Module,
                   f: FileItem) -> Path | None:
    """Where to save this file, or None if we already have it (or an identical copy)."""
    if manifest.is_current(module, f):
        return None
    path = download_path(cfg, course, module, f)
    key = manifest.key(module, f)
    if key not in manifest.data and path.exists() and path.stat().st_size == f.size:
        manifest.record(course, module, f, path)  # downloaded before the manifest existed
        return None
    if (dup := manifest.same_file_elsewhere(course, module, f)) is not None:
        manifest.record(course, module, f, manifest.root / manifest.data[dup]["path"], duplicate_of=dup)
        log.info("Skipping %s: already have it from another post", f.name)
        return None
    return path


def process_files(moodle: Moodle, digest: Digest, cfg: Config, want_text: bool) -> None:
    """For new or updated files: an AI summary of each (slides read as documents, so diagrams
    and equations count), plus a local copy when DOWNLOAD_FILES is on."""
    cache = SummaryCache(cfg.output_dir / "summary-cache.json")
    manifest = DownloadManifest(cfg) if cfg.download_files else None
    summarized = 0
    for course in digest.courses:
        for module in course.modules:
            for f in module.files:
                ext = Path(f.name).suffix.lower()
                wanted = bool(want_text and f.status and ext in SUMMARIZABLE)
                if wanted and (cached := cache.get(f)):
                    f.summary, wanted = cached, False
                wanted = wanted and summarized < MAX_SUMMARY_FILES
                target = None
                if manifest is not None and module.modname != "page" and not is_recording(f):
                    target = _plan_download(manifest, cfg, course, module, f)
                if not (wanted or target):
                    continue
                try:
                    data = moodle.download(f.download_url, MAX_DOWNLOAD_BYTES)
                except Exception as e:  # one bad file shouldn't sink the digest
                    log.warning("Download failed for %s: %s", f.name, e)
                    continue
                if target:
                    _save_download(manifest, course, module, f, target, data)
                if wanted:
                    summarized += 1
                    _summarize(cfg, course, module, f, data, cache)
    if manifest is not None:
        manifest.save()


def _save_download(manifest: DownloadManifest, course: Course, module: Module, f: FileItem,
                   target: Path, data: bytes) -> None:
    sha1 = hashlib.sha1(data).hexdigest()
    key = manifest.key(module, f)
    if (dup := manifest.same_content(sha1, exclude=key)) is not None:
        # Identical bytes under another name/post: point at the copy we have, don't store twice.
        manifest.record(course, module, f, manifest.root / manifest.data[dup]["path"], sha1, dup)
        log.info("Skipping %s: identical to %s", f.name, manifest.data[dup]["path"])
        return
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_bytes(data)
    manifest.record(course, module, f, target, sha1)
    log.info("Downloaded %s", target.relative_to(manifest.root))


def _summarize(cfg: Config, course: Course, module: Module, f: FileItem, data: bytes,
               cache: SummaryCache) -> None:
    from .summarize import summarize_file

    try:
        f.summary = summarize_file(cfg, course, module, f, data)
        if f.summary:
            cache.put(f, f.summary)
            return
    except Exception as e:
        log.warning("AI couldn't summarize %s: %s", f.name, e)
    _attach_text(f, data)  # fall back to raw text for the overall summary


def _attach_text(f: FileItem, data: bytes) -> None:
    try:
        text = extract_text(data, f.name)
    except Exception as e:
        log.warning("Couldn't read text from %s: %s", f.name, e)
        return
    text = re.sub(r"\n\s*\n+", "\n\n", text).strip()
    f.text = text[:EXCERPT_CHARS]
